
import openpyxl

# Load the original workbook (before conversion)
wb = openpyxl.load_workbook(r'd:\dke3\tooldke\code\extracted_data_final_v9.xlsx')

# Sheets to process (exclude 'img')
sheets_to_process = [sheet for sheet in wb.sheetnames if sheet != 'img']

for sheet_name in sheets_to_process:
    ws = wb[sheet_name]
    print(f"\nSheet {sheet_name} - Rows where Maximum = 2:")
    for row_idx in range(2, ws.max_row + 1):
        max_val = ws.cell(row=row_idx, column=4).value
        if max_val == 2:
            # Print the row data
            row_data = [ws.cell(row=row_idx, column=col).value for col in range(1, ws.max_column + 1)]
            print(f"Row {row_idx}: {row_data}")

wb.close()
