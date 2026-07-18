
import openpyxl
import os
from openpyxl.styles import PatternFill

# Configuration - use relative paths for portability
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
INPUT_FILE = os.path.join(SCRIPT_DIR, "code", "extracted_data_final_v9.xlsx")
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "code", "extracted_data_final_v9_with_averages.xlsx")

# Load the workbook
wb = openpyxl.load_workbook(INPUT_FILE)

# Define yellow fill
yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')

# Sheets to process (exclude 'img')
sheets_to_process = [sheet for sheet in wb.sheetnames if sheet != 'img']

for sheet_name in sheets_to_process:
    ws = wb[sheet_name]

    # Step 1: Convert Mean values (column E, index 5) from uA to mA if needed
    # Assume values > 10 are uA (since mA should be smaller)
    for row_idx in range(2, ws.max_row + 1):  # Start from row 2 (skip header)
        mean_cell = ws.cell(row=row_idx, column=5)
        if mean_cell.value is not None and isinstance(mean_cell.value, (int, float)):
            if mean_cell.value > 10:  # If value is large, it's uA → convert to mA
                mean_cell.value = mean_cell.value / 1000

    last_row = ws.max_row
    avg_row = last_row + 1

    ws.cell(row=avg_row, column=1, value='Average')

    # Apply yellow fill
    for col in range(1, ws.max_column + 1):
        ws.cell(row=avg_row, column=col).fill = yellow_fill

    # Add AVERAGE formulas
    # Column D (Maximum): keep
    col_d_letter = openpyxl.utils.get_column_letter(4)
    ws.cell(row=avg_row, column=4, value=f'=AVERAGE({col_d_letter}2:{col_d_letter}{last_row})')

    # Column C (Delta_t) → move to column E
    col_c_letter = openpyxl.utils.get_column_letter(3)
    ws.cell(row=avg_row, column=5, value=f'=AVERAGE({col_c_letter}2:{col_c_letter}{last_row})')

    # Column E (Mean) → move to column C
    col_e_letter = openpyxl.utils.get_column_letter(5)
    ws.cell(row=avg_row, column=3, value=f'=AVERAGE({col_e_letter}2:{col_e_letter}{last_row})')

# Save the modified workbook (leaving img sheet exactly as it is!)
wb.save(OUTPUT_FILE)
print("Done! Converted Mean to mA and added averages (img sheet left untouched)!")
