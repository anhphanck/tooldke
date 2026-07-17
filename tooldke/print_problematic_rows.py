
import openpyxl

wb = openpyxl.load_workbook(r'd:\dke3\tooldke\code\extracted_data_final_v9.xlsx')
sheets_to_process = [sheet for sheet in wb.sheetnames if sheet != 'img']

print("=== Problematic Rows ===")
for sheet_name in sheets_to_process:
    ws = wb[sheet_name]
    print(f"\n--- Sheet: {sheet_name} ---")
    for row_idx in range(2, ws.max_row + 1):
        max_val = ws.cell(row=row_idx, column=4).value
        if max_val == 2:
            filename = ws.cell(row=row_idx, column=2).value
            row_data = [ws.cell(row=row_idx, column=col).value for col in range(1, ws.max_column + 1)]
            print(f"Row {row_idx}: {row_data}")
            print(f"  Filename: {filename}")

wb.close()
