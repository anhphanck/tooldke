
import openpyxl

wb = openpyxl.load_workbook(r'd:\dke3\tooldke\code\extracted_data_final_v9.xlsx')
sheets_to_process = [sheet for sheet in wb.sheetnames if sheet != 'img']

with open(r'd:\dke3\tooldke\problematic_rows.txt', 'w', encoding='utf-8') as f:
    for sheet_name in sheets_to_process:
        ws = wb[sheet_name]
        f.write(f"=== Sheet: {sheet_name} ===\n")
        for row_idx in range(2, ws.max_row + 1):
            max_val = ws.cell(row=row_idx, column=4).value
            if max_val == 2:
                filename = ws.cell(row=row_idx, column=2).value
                row_data = [ws.cell(row=row_idx, column=col).value for col in range(1, ws.max_column + 1)]
                f.write(f"Row {row_idx}: {row_data}\n")
                f.write(f"  Filename: {filename}\n\n")

wb.close()
print("Wrote problematic rows to problematic_rows.txt")
