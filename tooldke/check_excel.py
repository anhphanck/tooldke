
import openpyxl

# Load the workbook
wb = openpyxl.load_workbook(r'd:\dke3\tooldke\code\extracted_data_final_v9.xlsx')

# List all sheet names
print("Sheet names:", wb.sheetnames)

# Let's check one sheet (e.g., T0) to see the structure
if 'T0' in wb.sheetnames:
    ws = wb['T0']
    print("\nT0 sheet - first 10 rows:")
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i > 10:
            break
        print(row)

wb.close()
