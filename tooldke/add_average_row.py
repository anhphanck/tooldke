
import openpyxl
from openpyxl.styles import PatternFill

# Load the workbook
wb = openpyxl.load_workbook(r'd:\dke3\tooldke\code\extracted_data_final_v9.xlsx')

# Define yellow fill
yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')

# Sheets to process (exclude 'img')
sheets_to_process = [sheet for sheet in wb.sheetnames if sheet != 'img']

for sheet_name in sheets_to_process:
    ws = wb[sheet_name]
    
    # Find the last row with data
    last_row = ws.max_row
    
    # Add average row below the last data row
    avg_row = last_row + 1
    
    # Set the first cell (maybe empty or "Average")
    ws.cell(row=avg_row, column=1, value='Average')
    
    # Apply yellow fill to the entire row
    for col in range(1, ws.max_column + 1):
        ws.cell(row=avg_row, column=col).fill = yellow_fill
    
    # Add AVERAGE formulas:
    # Column D (Maximum): keep
    col_d_letter = openpyxl.utils.get_column_letter(4)
    ws.cell(row=avg_row, column=4, value=f'=AVERAGE({col_d_letter}2:{col_d_letter}{last_row})')
    
    # Column C (Delta_t) → move to column E
    col_c_letter = openpyxl.utils.get_column_letter(3)
    ws.cell(row=avg_row, column=5, value=f'=AVERAGE({col_c_letter}2:{col_c_letter}{last_row})')
    
    # Column E (Mean) → move to column C
    col_e_letter = openpyxl.utils.get_column_letter(5)
    ws.cell(row=avg_row, column=3, value=f'=AVERAGE({col_e_letter}2:{col_e_letter}{last_row})')
    
    # Column F (RMS): leave empty

# Save the modified workbook
wb.save(r'd:\dke3\tooldke\code\extracted_data_final_v9_with_averages.xlsx')
print("Done! Averages updated. Saved as extracted_data_final_v9_with_averages.xlsx")
